# -*- coding: utf-8 -*-
##############################################################################
#
#  AuctionChamp - Professional Sports Auction Management Platform
#
#  Copyright (c) 2026 AuctionChamp.
#  All Rights Reserved.
#
#  CONFIDENTIAL & PROPRIETARY
#
#  This source code, including but not limited to its algorithms, business
#  logic, database structures, models, controllers, views, reports, templates,
#  APIs, documentation, and related materials, constitutes proprietary and
#  confidential information owned exclusively by AuctionChamp.
#
#  This software is protected by applicable copyright laws and international
#  intellectual property treaties. Unauthorized copying, reproduction,
#  modification, distribution, publication, sublicensing, reverse engineering,
#  decompilation, disassembly, disclosure, or use of this software, in whole
#  or in part, is strictly prohibited without the prior written permission of
#  AuctionChamp.
#
#  This software is licensed, not sold. Possession of the source code does not
#  grant any right to copy, modify, redistribute, or create derivative works
#  except as expressly permitted under a valid written license agreement with
#  AuctionChamp.
#
#  Any unauthorized use may result in civil and criminal penalties under
#  applicable intellectual property and copyright laws.
#
#  Company  : AuctionChamp
#  Website  : www.auctionchamp.live
#  Email    : auctionchamp.live@gmail.com
#
#  © 2026 AuctionChamp. All Rights Reserved.
#
##############################################################################

import io
import base64
import openpyxl
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, GradientFill
)
from openpyxl.utils import get_column_letter

from odoo import api, models, fields


# ─── Helpers ────────────────────────────────────────────────────────────────

def _side(style='thin', color='FFBDBDBD'):
    return Side(style=style, color=color)

BORDER_THIN = Border(
    left=_side(), right=_side(), top=_side(), bottom=_side()
)
BORDER_MED = Border(
    left=_side('medium', 'FF666666'), right=_side('medium', 'FF666666'),
    top=_side('medium', 'FF666666'), bottom=_side('medium', 'FF666666'),
)


# ─── Models ─────────────────────────────────────────────────────────────────

class AuctionPlanner(models.Model):
    _name = 'auction.planner'
    _description = 'Auction Planner'
    _rec_name = 'name'

    name = fields.Char(string='Plan Name', required=True)
    total_purse = fields.Integer(string='Total Purse', required=True, default=0)
    slab_ids = fields.One2many('auction.planner.slab', 'planner_id', string='Bid Slabs')
    player_ids = fields.One2many('auction.planner.player', 'planner_id', string='Players')

    # ── Computed summary ──────────────────────────────────────────────────────
    total_used = fields.Integer(
        string='Total Used',
        compute='_compute_summary', store=True,
    )
    remaining_balance = fields.Integer(
        string='Remaining Balance',
        compute='_compute_summary', store=True,
    )
    max_next_call = fields.Integer(
        string='Max Next Call',
        compute='_compute_summary', store=True,
    )

    @api.depends('total_purse', 'player_ids', 'player_ids.points', 'slab_ids',
                 'slab_ids.from_amount', 'slab_ids.increment')
    def _compute_summary(self):
        for rec in self:
            used = sum(p.points for p in rec.player_ids)
            remaining = rec.total_purse - used
            rec.total_used = used
            rec.remaining_balance = remaining
            rec.max_next_call = rec._snap_to_slab(remaining)

    def _snap_to_slab(self, amount):
        """Snap amount DOWN to the nearest valid slab step (mirrors auction_module logic)."""
        for slab in self.slab_ids.sorted('from_amount', reverse=True):
            if amount >= slab.from_amount:
                base = slab.from_amount
                inc = slab.increment
                if inc <= 0:
                    return base
                snapped = base + ((amount - base) // inc) * inc
                return int(min(snapped, amount))
        return amount

    # ── Excel export ──────────────────────────────────────────────────────────

    def action_export_excel(self):
        self.ensure_one()
        xlsx_data = self._build_excel()
        b64 = base64.b64encode(xlsx_data)
        filename = '{}.xlsx'.format(
            (self.name or 'AuctionPlan').replace('/', '_').replace('\\', '_')
        )
        attachment = self.env['ir.attachment'].create({
            'name': filename,
            'type': 'binary',
            'datas': b64,
            'mimetype': (
                'application/vnd.openxmlformats-officedocument'
                '.spreadsheetml.sheet'
            ),
            'res_model': self._name,
            'res_id': self.id,
        })
        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/{}?download=true'.format(attachment.id),
            'target': 'self',
        }

    # ── Excel builder ─────────────────────────────────────────────────────────

    def _build_excel(self):
        """Generate a two-sheet workbook:
        - "Planner"  : summary section + player table (with live formulas)
        - "Slabs"    : slab table referenced by the Max Next Call formula
        """
        wb = openpyxl.Workbook()

        # ── Sheet 1: Planner ──────────────────────────────────────────────
        ws = wb.active
        ws.title = 'Planner'

        # column widths
        ws.column_dimensions['A'].width = 24
        ws.column_dimensions['B'].width = 36
        ws.column_dimensions['C'].width = 18

        # ── Title ─────────────────────────────────────────────────────────
        self._xl_title(ws, 'A1', 'B1', 'Auction Planner')

        # ── Plan name row ─────────────────────────────────────────────────
        self._xl_label(ws, 'A2', 'Plan')
        ws['B2'] = self.name or ''
        ws['B2'].font = Font(bold=True, size=11, color='FF1F4E79')
        ws['B2'].alignment = Alignment(horizontal='left', vertical='center')

        # ── SUMMARY section ───────────────────────────────────────────────
        #   Row 4: section header
        #   Row 5: Total Purse           B5  ← hard value
        #   Row 6: Total Used            B6  = SUM(C12:C10000)
        #   Row 7: Remaining Balance     B7  = B5 - B6
        #   Row 8: Max Next Call         B8  = slab-snap formula

        self._xl_section_header(ws, 'A4', 'B4', 'SUMMARY')

        PURSE_ROW   = 5
        USED_ROW    = 6
        REM_ROW     = 7
        MAX_ROW     = 8
        HDR_ROW     = 11   # player table header
        DATA_START  = 12   # first player data row

        summary_rows = [
            (PURSE_ROW, 'Total Purse',       self.total_purse,       None),
            (USED_ROW,  'Total Used',         None,                   f'=SUM(C{DATA_START}:C10000)'),
            (REM_ROW,   'Remaining Balance',  None,                   f'=B{PURSE_ROW}-B{USED_ROW}'),
            (MAX_ROW,   'Max Next Call',      None,                   self._max_next_call_formula(REM_ROW)),
        ]

        FILL_PURSE   = PatternFill('solid', fgColor='FFD6E4F7')
        FILL_USED    = PatternFill('solid', fgColor='FFFFF2CC')
        FILL_REM     = PatternFill('solid', fgColor='FFE2EFDA')
        FILL_MAX     = PatternFill('solid', fgColor='FFFCE4D6')
        fills = [FILL_PURSE, FILL_USED, FILL_REM, FILL_MAX]

        for i, (row, label, val, formula) in enumerate(summary_rows):
            a_cell = ws.cell(row=row, column=1, value=label)
            a_cell.font = Font(bold=True, size=10, color='FF333333')
            a_cell.fill = fills[i]
            a_cell.border = BORDER_THIN
            a_cell.alignment = Alignment(horizontal='left', vertical='center')

            b_cell = ws.cell(row=row, column=2)
            b_cell.value = formula if formula else val
            b_cell.font = Font(bold=True, size=11, color='FF1F4E79')
            b_cell.fill = fills[i]
            b_cell.border = BORDER_THIN
            b_cell.alignment = Alignment(horizontal='right', vertical='center')

            # number format
            b_cell.number_format = '#,##0'

        ws.row_dimensions[PURSE_ROW].height = 22
        ws.row_dimensions[USED_ROW].height  = 22
        ws.row_dimensions[REM_ROW].height   = 22
        ws.row_dimensions[MAX_ROW].height   = 22

        # ── Player table ──────────────────────────────────────────────────
        self._xl_section_header(ws, 'A10', 'C10', 'PLAYERS')

        # Header row
        HDRS = ['#', 'Player Name', 'Points']
        HDRC = ['FFD9E1F2', 'FFD9E1F2', 'FFD9E1F2']
        for col, (hdr, clr) in enumerate(zip(HDRS, HDRC), start=1):
            c = ws.cell(row=HDR_ROW, column=col, value=hdr)
            c.font = Font(bold=True, size=10, color='FF1F4E79')
            c.fill = PatternFill('solid', fgColor=clr)
            c.border = BORDER_MED
            c.alignment = Alignment(horizontal='center', vertical='center')
        ws.row_dimensions[HDR_ROW].height = 20

        # Data rows
        players = self.player_ids
        for idx, player in enumerate(players, start=1):
            row = DATA_START + idx - 1
            ws.cell(row=row, column=1, value=idx).alignment = Alignment(horizontal='center')
            ws.cell(row=row, column=2, value=player.player_name)
            pts_cell = ws.cell(row=row, column=3, value=player.points)
            pts_cell.number_format = '#,##0'
            for col in range(1, 4):
                c = ws.cell(row=row, column=col)
                c.border = BORDER_THIN
                fill_color = 'FFFFFFFF' if idx % 2 == 0 else 'FFF8F9FF'
                c.fill = PatternFill('solid', fgColor=fill_color)
            ws.row_dimensions[row].height = 18

        # Freeze panes below summary + header
        ws.freeze_panes = ws.cell(row=DATA_START, column=1)

        # ── Sheet 2: Slabs ────────────────────────────────────────────────
        ws2 = wb.create_sheet('Slabs')
        ws2.column_dimensions['A'].width = 16
        ws2.column_dimensions['B'].width = 16
        ws2.column_dimensions['C'].width = 16

        slab_hdrs = ['From Amount', 'To Amount', 'Increment']
        for col, hdr in enumerate(slab_hdrs, start=1):
            c = ws2.cell(row=1, column=col, value=hdr)
            c.font = Font(bold=True, color='FFFFFFFF')
            c.fill = PatternFill('solid', fgColor='FF2E4057')
            c.border = BORDER_THIN
            c.alignment = Alignment(horizontal='center')

        slabs_sorted = self.slab_ids.sorted('from_amount')
        for row_i, slab in enumerate(slabs_sorted, start=2):
            for col, val in enumerate(
                [slab.from_amount, slab.to_amount, slab.increment], start=1
            ):
                c = ws2.cell(row=row_i, column=col, value=val)
                c.border = BORDER_THIN
                c.number_format = '#,##0'
                fill_color = 'FFFFFFFF' if row_i % 2 == 0 else 'FFEFF3FA'
                c.fill = PatternFill('solid', fgColor=fill_color)

        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()

    @staticmethod
    def _max_next_call_formula(rem_row):
        """
        Excel formula that replicates _snap_to_slab():
          - MATCH(rem, sorted from_amount, 1)  → finds the applicable slab row
          - base + FLOOR(rem - base, increment) → snaps down to nearest step
          - IFERROR fallback returns rem as-is if no slab matches
        Slabs sheet columns: A=from_amount, B=to_amount, C=increment (data from row 2)
        MATCH type 1 requires from_amount sorted ASCENDING (we sort when writing).
        """
        rem = f'B{rem_row}'
        a = "Slabs!$A$2:$A$1000"
        c = "Slabs!$C$2:$C$1000"
        match = f'MATCH({rem},{a},1)'
        base  = f'INDEX({a},{match})'
        inc   = f'INDEX({c},{match})'
        snapped = f'{base}+FLOOR({rem}-{base},{inc})'
        return f'=IFERROR({snapped},{rem})'

    # ── Style helpers ─────────────────────────────────────────────────────────

    @staticmethod
    def _xl_title(ws, start_cell, end_cell, text):
        c = ws[start_cell]
        c.value = text
        c.font = Font(bold=True, size=14, color='FFFFFFFF')
        c.fill = PatternFill('solid', fgColor='FF1F4E79')
        c.alignment = Alignment(horizontal='center', vertical='center')
        c.border = BORDER_MED
        ws.merge_cells(f'{start_cell}:{end_cell}')
        ws.row_dimensions[int(start_cell[1:])].height = 28

    @staticmethod
    def _xl_label(ws, cell_ref, text):
        c = ws[cell_ref]
        c.value = text
        c.font = Font(bold=True, size=10, color='FF555555')
        c.alignment = Alignment(horizontal='left', vertical='center')

    @staticmethod
    def _xl_section_header(ws, start_cell, end_cell, text):
        col = start_cell[0]
        row = int(start_cell[1:])
        c = ws[start_cell]
        c.value = text
        c.font = Font(bold=True, size=10, color='FFFFFFFF', italic=True)
        c.fill = PatternFill('solid', fgColor='FF2E4057')
        c.alignment = Alignment(horizontal='left', vertical='center',
                                indent=1)
        ws.merge_cells(f'{start_cell}:{end_cell}')
        ws.row_dimensions[row].height = 20


class AuctionPlannerSlab(models.Model):
    _name = 'auction.planner.slab'
    _description = 'Auction Planner Bid Slab'
    _order = 'from_amount asc'

    planner_id = fields.Many2one('auction.planner', ondelete='cascade')
    from_amount = fields.Integer(string='From Amount', required=True)
    to_amount   = fields.Integer(string='To Amount',   required=True)
    increment   = fields.Integer(string='Increment',   required=True)


class AuctionPlannerPlayer(models.Model):
    _name = 'auction.planner.player'
    _description = 'Auction Planner Player'
    _order = 'sequence, id'

    planner_id  = fields.Many2one('auction.planner', ondelete='cascade')
    sequence    = fields.Integer(default=10)
    player_name = fields.Char(string='Player Name', required=True)
    points      = fields.Integer(string='Points', default=0)
