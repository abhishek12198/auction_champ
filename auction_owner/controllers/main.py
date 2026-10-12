# -*- coding: utf-8 -*-
##############################################################################
#
#  AuctionChamp - Professional Sports Auction Management Platform
#
#  Copyright (c) 2026 AuctionChamp.
#  All Rights Reserved.
#
#  CONFIDENTIAL & PROPRIETARY
#
#  This source code, including but not limited to its algorithms, business
#  logic, database structures, models, controllers, views, reports, templates,
#  APIs, documentation, and related materials, constitutes proprietary and
#  confidential information owned exclusively by AuctionChamp.
#
#  This software is protected by applicable copyright laws and international
#  intellectual property treaties. Unauthorized copying, reproduction,
#  modification, distribution, publication, sublicensing, reverse engineering,
#  decompilation, disassembly, disclosure, or use of this software, in whole
#  or in part, is strictly prohibited without the prior written permission of
#  AuctionChamp.
#
#  This software is licensed, not sold. Possession of the source code does not
#  grant any right to copy, modify, redistribute, or create derivative works
#  except as expressly permitted under a valid written license agreement with
#  AuctionChamp.
#
#  Any unauthorized use may result in civil and criminal penalties under
#  applicable intellectual property and copyright laws.
#
#  Company  : AuctionChamp
#  Website  : www.auctionchamp.live
#  Email    : auctionchamp.live@gmail.com
#
#  © 2026 AuctionChamp. All Rights Reserved.
#
##############################################################################

import json
import logging

from odoo import SUPERUSER_ID, fields as odoo_fields, http
from odoo.http import request
from werkzeug.urls import url_quote
from odoo.addons.web.controllers.main import Home

_logger = logging.getLogger(__name__)


# ── Login redirect: send Owner users straight to the Owner Console ─────────────
class AuctionOwnerHome(Home):
    @http.route('/web', type='http', auth='none')
    def web_client(self, s_action=None, **kw):
        if request.session.uid:
            try:
                user = request.env['res.users'].sudo().browse(request.session.uid)
                if user.has_group('auction_owner.group_auction_owner'):
                    return request.redirect('/auction/owner/console', 303)
            except Exception:
                pass
        return super().web_client(s_action=s_action, **kw)


class AuctionOwnerController(http.Controller):

    # ── Helpers ────────────────────────────────────────────────────────────

    @staticmethod
    def _resolve_tournament():
        """Return the tournament for the current user.

        Team owners must always see the tournament of their assigned team,
        never a random active tournament or a leftover Active Tournament.
        """
        user = request.env['res.users'].sudo().browse(request.session.uid)
        team = user.auction_team_id or getattr(user, 'team_id', False)
        if team and team.tournament_id:
            return team.tournament_id
        if user.has_group('auction_owner.group_auction_owner') and not user.has_group(
            'auction_module.group_auction_group_admin'
        ):
            return team.tournament_id if team else False
        if user.tournament_id:
            return user.tournament_id
        return request.env['auction.tournament'].sudo().search(
            [('active', '=', True)], limit=1
        )

    @staticmethod
    def _tournament_from_request(**kw):
        """Resolve tournament by slug (?t=) so two live auctions do not mix."""
        slug = (kw.get('t') or '').strip()
        if slug:
            found = request.env['auction.tournament'].sudo().search(
                [('slug', '=', slug)], limit=1
            )
            if found:
                return found
        return None

    @staticmethod
    def _pub_img(model, record_id, field):
        return '/auction/public/image/%s/%d/%s' % (model, record_id, field)

    @staticmethod
    def _fmt_pts(tournament, amount):
        if tournament and hasattr(tournament, 'format_points'):
            return tournament.format_points(amount)
        return '{:,}'.format(int(amount or 0))

    @staticmethod
    def _get_effective_base(auction, player):
        if player and player.tier_id and auction.tier_limit_ids:
            tl = auction.tier_limit_ids.filtered(
                lambda l: l.tier_id.id == player.tier_id.id
            )
            if tl and tl[0].base_point > 0:
                return tl[0].base_point
        return auction.base_point

    @staticmethod
    def _next_bid(auction, current_bid, effective_base):
        if current_bid == 0:
            return effective_base
        slabs = auction.auction_bid_slab_ids.sorted('from_amount', reverse=True)
        for slab in slabs:
            if current_bid >= slab.from_amount:
                return current_bid + slab.increment
        return current_bid + 1

    def _squad_for_auction(self, auction):
        """Return a list of player dicts for the given auction record."""
        env = request.env
        ap_records = env['auction.auction.player'].sudo().search(
            [('auction_id', '=', auction.id)]
        )
        squad = []
        for ap in ap_records:
            p = ap.player_id
            squad.append({
                'id': p.id,
                'name': p.name or '',
                'photo_url': self._pub_img('auction.team.player', p.id, 'photo') if p.photo else '',
                'role': p.role or '',
                'tier_name': p.tier_id.name if p.tier_id else '',
                'tier_color': p.tier_id.color if p.tier_id else '#3498db',
                'points': ap.points or 0,
            })
        return squad

    def _team_payload(self, auc, player, include_squad=False):
        """Build a team data dict for the given auction record."""
        team = auc.team_id
        effective_base = self._get_effective_base(auc, player)
        current_bid_val = (player.current_bid or 0) if player else 0
        next_bid_val = self._next_bid(auc, current_bid_val, effective_base)
        # Pass player so tier max_call cap and per-tier base reserves are applied
        live_max_call = auc.get_max_bid_for_team(auc, player)

        # Resolve the minimum bid required for the current player's tier
        _tier_base = 0
        _tier_name = ''
        if player and player.tier_id:
            _tier_name = player.tier_id.name or ''
            if auc.tier_limit_ids:
                _tl = auc.tier_limit_ids.filtered(lambda l: l.tier_id.id == player.tier_id.id)
                _tier_base = (_tl[0].base_point if _tl and _tl[0].base_point > 0
                              else (auc.base_point or 0))
            else:
                _tier_base = auc.base_point or 0

        can_bid = True
        can_bid_reason = ''

        if not player:
            can_bid = False
            can_bid_reason = 'No player on stage'
        elif player.state != 'auction':
            can_bid = False
            can_bid_reason = 'Waiting for the next player'
        elif auc.remaining_players_count <= 0:
            can_bid = False
            can_bid_reason = 'Squad full'
        elif auc.remaining_points <= 0:
            can_bid = False
            can_bid_reason = 'No budget left'
        elif _tier_base > 0 and auc.remaining_points < _tier_base:
            can_bid = False
            can_bid_reason = 'Purse below %s minimum (%s)' % (
                _tier_name, self._fmt_pts(auc.tournament_id, _tier_base),
            )
        elif live_max_call <= 0:
            can_bid = False
            can_bid_reason = 'Budget reserved for other players'
        elif next_bid_val > live_max_call:
            can_bid = False
            can_bid_reason = 'Max call: %s' % self._fmt_pts(auc.tournament_id, live_max_call)

        # Current-bid restriction: owner cannot raise their own bid
        if can_bid and player and player.current_bid_team_id and \
                player.current_bid_team_id.id == team.id:
            can_bid = False
            can_bid_reason = 'You have the current bid — wait for a rival'

        bid_timer = (
            player.get_owner_bid_timer_state()
            if player and hasattr(player, 'get_owner_bid_timer_state')
            else {'frozen': False}
        )
        is_leading_bid = bool(
            player and player.current_bid_team_id and player.current_bid_team_id.id == team.id
        )
        if bid_timer.get('frozen') and not is_leading_bid:
            can_bid = False
            can_bid_reason = 'Time up — only the last bid stays open'

        if can_bid and player and player.tier_id and auc.tier_limit_ids:
            tl = auc.tier_limit_ids.filtered(lambda l: l.tier_id.id == player.tier_id.id)
            if tl:
                already_sold = request.env['auction.auction.player'].sudo().search_count([
                    ('auction_id', '=', auc.id),
                    ('player_id.tier_id', '=', player.tier_id.id),
                ])
                if already_sold >= tl[0].max_players:
                    can_bid = False
                    can_bid_reason = 'Tier slot full'

        payload = {
            'id': team.id,
            'auction_id': auc.id,
            'name': team.name or '',
            'logo_url': self._pub_img('auction.team', team.id, 'logo') if team.logo else '',
            'remaining_points': auc.remaining_points,
            'total_points': auc.total_point,
            'max_call': live_max_call,
            'next_bid': next_bid_val,
            'effective_base': effective_base,
            'can_bid': can_bid,
            'can_bid_reason': can_bid_reason,
            'is_leading': bool(
                player and player.current_bid_team_id and
                player.current_bid_team_id.id == team.id
            ),
            'player_count': len(auc.player_ids),
            'max_players': auc.max_players,
            'manager': team.manager or '',
        }
        if include_squad:
            payload['squad'] = self._squad_for_auction(auc)
        return payload

    # ── Owner Console page ─────────────────────────────────────────────────

    @http.route('/auction/owner/console', type='http', auth='user', website=False)
    def owner_console(self, **kw):
        user = request.env.user
        if not user.has_group('auction_owner.group_auction_owner'):
            return request.not_found()

        tournament = self._resolve_tournament()
        return request.render(
            'auction_owner.owner_console_template',
            {
                'tournament': tournament,
                'user': user,
                'favicon_url': '/web/image/res.company/%d/favicon' % request.env.company.id,
                'db_name': request.env.cr.dbname,
            },
        )

    # ── My Squad: roster PDF (existing auction.auction report) ─────────────

    @http.route('/auction/owner/my-squad/roster.pdf', type='http', auth='user', website=False)
    def owner_squad_roster_pdf(self, **kw):
        """Squad roster PDF for the signed-in owner's own team only."""
        user = request.env.user
        if not user.has_group('auction_owner.group_auction_owner'):
            return request.not_found()
        team = user.auction_team_id
        tournament = self._resolve_tournament()
        if not team or not tournament:
            return request.not_found()
        auction = request.env['auction.auction'].sudo().search([
            ('team_id', '=', team.id),
            ('tournament_id', '=', tournament.id),
        ], limit=1)
        if not auction:
            return request.not_found()
        # Render as the superuser: the roster template reads models (venue, tiers...)
        # that owner users have no ACL for. Access was already restricted above to
        # the owner's own team auction, and only that id is rendered.
        report = request.env.ref('auction_module.action_report_auction_players').with_user(SUPERUSER_ID)
        pdf, _fmt = report._render_qweb_pdf(auction.ids)
        filename = '%s Squad Roster.pdf' % (team.name or 'Team')
        return request.make_response(pdf, headers=[
            ('Content-Type', 'application/pdf'),
            ('Content-Length', len(pdf)),
            ('Content-Disposition', "inline; filename*=UTF-8''%s" % url_quote(filename, safe='')),
            ('Cache-Control', 'no-store'),
        ])

    @staticmethod
    def _owner_player_profile(pl, sport, fallback_role=''):
        """Role + attribute rows for the players board, specific to each sport."""
        from odoo.addons.auction_module.services.auction_live_payload import (
            _balance_player_profile,
        )
        if sport == 'badminton':
            lb = pl.badminton_card_labels() or {}
            attrs = []
            for key, label in (('category', 'Category'), ('style', 'Style'),
                               ('level', 'Level'), ('hand', 'Hand')):
                if lb.get(key):
                    attrs.append({'k': label, 'v': str(lb[key])})
            if lb.get('age'):
                attrs.append({'k': 'Age', 'v': str(lb['age'])})
            strengths = ', '.join(
                ('%s %s' % (x.get('icon') or '', x.get('name') or '')).strip()
                for x in (lb.get('strengths') or []) if x.get('name'))
            if strengths:
                attrs.append({'k': 'Strengths', 'v': strengths})
            return (lb.get('category') or pl.role or fallback_role or ''), attrs
        if sport == 'kabaddi':
            attrs = []
            if getattr(pl, 'age', 0):
                attrs.append({'k': 'Age', 'v': str(pl.age)})
            return (pl.kabaddi_role_label() or pl.role or fallback_role or ''), attrs
        role, attrs = _balance_player_profile(pl, False, sport == 'football')
        return (role or fallback_role or ''), attrs

    # ── Players board (Sold / Unsold / In auction) ─────────────────────────

    @http.route('/auction/owner/players', type='http', auth='user', website=False)
    def owner_players(self, bucket='sold', offset=0, limit=20, q='', **kw):
        """Paginated player list for the owner's own tournament (bal snapshot)."""
        user = request.env.user
        if not user.has_group('auction_owner.group_auction_owner'):
            return request.not_found()
        tournament = self._resolve_tournament()
        if not tournament:
            return request.not_found()
        from odoo.addons.auction_module.services.auction_live_snapshot_service import (
            get_or_rebuild_snapshot,
        )
        from odoo.addons.auction_module.services.auction_live_payload import (
            slice_balance_players,
        )
        snap = get_or_rebuild_snapshot(request.env, tournament, 'bal') or {}
        payload = slice_balance_players(
            snap, bucket, offset=offset, limit=min(int(limit or 20), 50) if str(limit).isdigit() else 20,
            query=q,
        )
        # The snapshot carries the raw per-player base_price (usually 0). Use the same
        # effective base the live bid panel uses (tier base, else auction base).
        rows = payload['players']
        sport = tournament.tournament_type or 'cricket'
        if rows:
            players = request.env['auction.team.player'].sudo().browse(
                [r['id'] for r in rows]).exists()
            by_id = {pl.id: pl for pl in players}
            auctions = request.env['auction.auction'].sudo().search(
                [('tournament_id', '=', tournament.id)])
            for row in rows:
                pl = by_id.get(row['id'])
                base = 0
                for auc in auctions:
                    base = max(base, self._get_effective_base(auc, pl))
                row['base_price'] = base or row.get('base_price') or 0
                # Full sport-specific attribute set (cricket / football / badminton / kabaddi).
                if pl and row.get('name') != '???':
                    row['role'], row['attrs'] = self._owner_player_profile(pl, sport, row.get('role'))
        # Never hand player phone numbers to team owners.
        payload['players'] = [
            {k: v for k, v in row.items() if k != 'contact'} for row in payload['players']
        ]
        payload['my_team_id'] = user.auction_team_id.id if user.auction_team_id else 0
        return request.make_response(json.dumps(payload), [
            ('Content-Type', 'application/json'),
            ('Cache-Control', 'no-store'),
        ])

    # ── Owner Console data endpoint ────────────────────────────────────────

    @http.route('/auction/owner/data', type='http', auth='user', website=False, csrf=False)
    def owner_data(self, **kw):
        user = request.env.user.sudo()
        env = request.env

        tournament = self._resolve_tournament()

        result = {
            'tournament': {},
            'current_player': None,
            'my_team': None,
            'teams': [],
            'slabs': [],
            'presets': [],
            'owner': {'name': user.name or '', 'login': user.login or ''},
        }

        if tournament:
            result['tournament'] = {
                'name': tournament.name or '',
                'theme': tournament.player_display_template or 'vanilla',
                'logo_url': self._pub_img('auction.tournament', tournament.id, 'logo') if tournament.logo else '',
                'slug': tournament.slug or '',
                'db_name': request.env.cr.dbname,
                'point_unit': (
                    tournament.get_point_unit_js()
                    if hasattr(tournament, 'get_point_unit_js') else None
                ),
            }
            if tournament.preset_points:
                try:
                    result['presets'] = [
                        int(x.strip())
                        for x in tournament.preset_points.split(',')
                        if x.strip().lstrip('-').isdigit()
                    ]
                except Exception:
                    pass

        # ── Current player on stage ────────────────────────────────────────
        on_stage_domain = [('is_on_stage', '=', True)]
        if tournament:
            on_stage_domain.append(('tournament_id', '=', tournament.id))
        current_player = env['auction.team.player'].sudo().search(
            on_stage_domain, limit=1
        ) if tournament else env['auction.team.player'].sudo().browse()

        if current_player:
            base_price = 0
            auc_domain = [('tournament_id', '=', tournament.id)] if tournament else []
            auctions_all = env['auction.auction'].sudo().search(auc_domain)
            for auc in auctions_all:
                base = self._get_effective_base(auc, current_player)
                if base > base_price:
                    base_price = base

            current_bid_team = None
            if current_player.current_bid_team_id:
                t = current_player.current_bid_team_id
                current_bid_team = {
                    'id': t.id,
                    'name': t.name or '',
                    'logo_url': self._pub_img('auction.team', t.id, 'logo') if t.logo else '',
                }
            sold_team = None
            sold_points = 0
            if current_player.state == 'sold':
                buyer = current_player.assigned_team_id or current_player.current_bid_team_id
                if buyer:
                    sold_team = {
                        'id': buyer.id,
                        'name': buyer.name or '',
                        'logo_url': self._pub_img('auction.team', buyer.id, 'logo') if buyer.logo else '',
                    }
                sale = env['auction.auction.player'].sudo().search(
                    [('player_id', '=', current_player.id)], order='id desc', limit=1)
                sold_points = int(sale.points or 0) if sale else int(current_player.current_bid or 0)

            result['current_player'] = {
                'id': current_player.id,
                'name': current_player.name or '',
                'photo_url': self._pub_img('auction.team.player', current_player.id, 'photo') if current_player.photo else '',
                'role': current_player.role or '',
                'tier_name': current_player.tier_id.name if current_player.tier_id else '',
                'tier_color': current_player.tier_color or '#2252b5',
                'state': current_player.state,
                'sl_no': current_player.sl_no or 0,
                'base_price': base_price,
                'current_bid': current_player.current_bid or 0,
                'current_bid_team': current_bid_team,
                'sold_team': sold_team,
                'sold_points': sold_points,
                'batting_style': current_player.batting_style or '',
                'bowling_style': current_player.bowling_style or '',
                'bid_timer': current_player.get_owner_bid_timer_state(),
            }

        # ── Teams ─────────────────────────────────────────────────────────
        auc_domain = [('tournament_id', '=', tournament.id)] if tournament else []
        auctions = env['auction.auction'].sudo().search(auc_domain)

        my_team_id = user.auction_team_id.id if user.auction_team_id else False

        for auc in auctions:
            if not auc.team_id:
                continue
            is_mine = (auc.team_id.id == my_team_id)
            payload = self._team_payload(auc, current_player, include_squad=True)
            if is_mine:
                result['my_team'] = payload
                # Expose this auction's bid slabs so the frontend can compute
                # slab-aware step increments for the +/- bid buttons.
                result['slabs'] = [
                    {
                        'from_amount': s.from_amount,
                        'to_amount': s.to_amount,
                        'increment': s.increment,
                    }
                    for s in auc.auction_bid_slab_ids.sorted('from_amount', reverse=True)
                ]
            result['teams'].append(payload)

        # Sold / Unsold / In-auction counts for the header chips (cached bal snapshot,
        # the same one Bid Summary reads; no extra polling endpoint).
        result['player_counts'] = {'sold': 0, 'unsold': 0, 'auction': 0}
        if tournament:
            try:
                from odoo.addons.auction_module.services.auction_live_snapshot_service import (
                    get_or_rebuild_snapshot,
                )
                snap = get_or_rebuild_snapshot(env, tournament, 'bal') or {}
                result['player_counts'].update(snap.get('player_counts') or {})
            except Exception:
                _logger.exception('owner_data: player counts failed')

        result['counter'] = {
            'started_at': (
                tournament.counter_started_at.isoformat()
                if tournament and tournament.counter_started_at else None
            ),
            'hammer_count': tournament.hammer_count if tournament else 3,
            'age_seconds': (
                (odoo_fields.Datetime.now() - tournament.counter_started_at).total_seconds()
                if tournament and tournament.counter_started_at else None
            ),
        }

        # ── Revoke settings ────────────────────────────────────────────────
        my_team_id = user.auction_team_id.id if user.auction_team_id else -1
        revoke_feature_on = bool(
            tournament and tournament.revoke_enabled and tournament.max_revokes > 0
        )
        revokes_remaining = (
            max(0, tournament.max_revokes - tournament.revokes_used)
            if revoke_feature_on else 0
        )
        # Owner can revoke only if they are the current leading bidder
        is_leading = bool(
            current_player and
            current_player.current_bid_team_id and
            current_player.current_bid_team_id.id == my_team_id
        )
        result['revoke'] = {
            'feature_on':  revoke_feature_on,
            'remaining':   revokes_remaining,
            'max':         tournament.max_revokes if tournament else 0,
            'used':        tournament.revokes_used if tournament else 0,
            'can_revoke':  revoke_feature_on and revokes_remaining > 0 and is_leading,
        }

        return request.make_response(
            json.dumps(result),
            headers=[
                ('Content-Type', 'application/json'),
                ('Cache-Control', 'no-store'),
            ],
        )

    # ── Counter check (polled by owner consoles every 2 s) ─────────────────

    @http.route('/auction/owner/counter-check', type='http', auth='public',
                website=False, csrf=False, methods=['GET'])
    def counter_check(self, **kw):
        tournament = self._tournament_from_request(**kw)
        if not tournament:
            tournament = request.env['auction.tournament'].sudo().search(
                [('active', '=', True)], limit=1
            )
        result = {'started_at': None, 'hammer_count': 3, 'age_seconds': None}
        if tournament and tournament.counter_started_at:
            result['started_at']   = tournament.counter_started_at.isoformat()
            result['hammer_count'] = tournament.hammer_count or 3
            result['age_seconds']  = (
                odoo_fields.Datetime.now() - tournament.counter_started_at
            ).total_seconds()
        return request.make_response(
            json.dumps(result),
            headers=[
                ('Content-Type', 'application/json'),
                ('Cache-Control', 'no-store'),
            ],
        )

    # ── Live Bid (public, polled by display page) ──────────────────────────

    @http.route('/auction/owner/live-bid', type='http', auth='public',
                website=False, csrf=False, methods=['GET'])
    def live_bid(self, **kw):
        tournament = self._tournament_from_request(**kw)
        domain = [('is_on_stage', '=', True)]
        if tournament:
            domain.append(('tournament_id', '=', tournament.id))
        player = request.env['auction.team.player'].sudo().search(
            domain, limit=1
        )
        result = {
            'has_bid': False,
            'player_id': None,
            'current_bid': 0,
            'team_id': None,
            'team_name': None,
            'team_logo_url': None,
        }
        if player:
            result['player_id'] = player.id
            if player.current_bid and player.current_bid_team_id:
                t = player.current_bid_team_id
                result.update({
                    'has_bid': True,
                    'current_bid': player.current_bid,
                    'team_id': t.id,
                    'team_name': t.name or '',
                    'team_logo_url': self._pub_img('auction.team', t.id, 'logo') if t.logo else '',
                })
        return request.make_response(
            json.dumps(result),
            headers=[
                ('Content-Type', 'application/json'),
                ('Cache-Control', 'no-store'),
            ],
        )

    # ── Enable Counter ──────────────────────────────────────────────────────

    @http.route('/auction/owner/enable-counter', type='http', auth='public',
                website=False, csrf=False, methods=['GET', 'POST'])
    def enable_counter(self, **kw):
        tournament = self._tournament_from_request(**kw)
        if not tournament:
            tournament = request.env['auction.tournament'].sudo().search(
                [('active', '=', True)], limit=1
            )
        if not tournament:
            _logger.warning('enable_counter: no active tournament found')
            payload = {'ok': False, 'reason': 'no_tournament'}
        else:
            try:
                tournament.sudo().write({'counter_started_at': odoo_fields.Datetime.now()})
                payload = {'ok': True, 'ts': odoo_fields.Datetime.now().isoformat()}
            except Exception as e:
                _logger.error('enable_counter: write failed – %s', e)
                payload = {'ok': False, 'reason': 'write_error', 'detail': str(e)}
        return request.make_response(
            json.dumps(payload),
            headers=[
                ('Content-Type', 'application/json'),
                ('Cache-Control', 'no-store'),
                ('Access-Control-Allow-Origin', '*'),
            ],
        )

    # ── Place bid ──────────────────────────────────────────────────────────

    @http.route('/auction/owner/place-bid', type='json', auth='user', website=False, csrf=False)
    def place_bid(self, player_id, team_id, bid_amount, **kw):
        user = request.env.user.sudo()
        env = request.env

        # Owners can only bid for their own assigned team
        if not user.auction_team_id or user.auction_team_id.id != int(team_id):
            return {'success': False, 'error': 'You can only bid for your own team.'}

        player = env['auction.team.player'].sudo().browse(int(player_id))
        if not player.exists() or not player.is_on_stage:
            return {'success': False, 'error': 'Player is not currently on stage.'}
        if player.state != 'auction':
            return {'success': False, 'error': 'This player is closed. Waiting for the next player.'}

        # Tournament isolation: player must belong to the user's assigned tournament
        tournament = self._resolve_tournament()
        if tournament and player.tournament_id and player.tournament_id.id != tournament.id:
            return {'success': False, 'error': 'Player does not belong to your tournament.'}

        if player.current_bid_team_id and player.current_bid_team_id.id == int(team_id):
            return {'success': False, 'error': 'You already have the highest bid. Wait for another team to bid.'}

        # Lock the player row so two owners cannot both pass an expiring window.
        env.cr.execute(
            'SELECT id FROM auction_team_player WHERE id = %s FOR UPDATE',
            (player.id,),
        )
        player.invalidate_cache()
        if player._owner_bid_window_closed(player):
            return {
                'success': False,
                'error': 'Time is up. Only the last bid stays open.',
            }

        auction = env['auction.auction'].sudo().search(
            [('team_id', '=', int(team_id))], limit=1
        )
        if not auction:
            return {'success': False, 'error': 'Auction record not found for your team.'}

        bid_amount = int(bid_amount)
        effective_base = self._get_effective_base(auction, player)

        if bid_amount < effective_base:
            return {
                'success': False,
                'error': 'Bid must be at least %s (base price).' % self._fmt_pts(tournament, effective_base),
            }

        # Guard: purse must cover the tier's minimum before going further
        if player.tier_id and auction.tier_limit_ids:
            _tl_base = auction.tier_limit_ids.filtered(lambda l: l.tier_id.id == player.tier_id.id)
            _tier_min = (_tl_base[0].base_point if _tl_base and _tl_base[0].base_point > 0
                         else (auction.base_point or 0))
            if _tier_min > 0 and auction.remaining_points < _tier_min:
                return {
                    'success': False,
                    'error': 'Insufficient purse for "%s" tier (requires %s, you have %s).' % (
                        player.tier_id.name,
                        self._fmt_pts(tournament, _tier_min),
                        self._fmt_pts(tournament, auction.remaining_points),
                    ),
                }

        # Pass player so tier max_call cap and per-tier budget reserves are applied
        live_max_call = auction.get_max_bid_for_team(auction, player)
        if bid_amount > live_max_call:
            return {
                'success': False,
                'error': 'Bid of %s exceeds your max call of %s.' % (
                    self._fmt_pts(tournament, bid_amount),
                    self._fmt_pts(tournament, live_max_call),
                ),
            }

        if player.tier_id and auction.tier_limit_ids:
            tl = auction.tier_limit_ids.filtered(lambda l: l.tier_id.id == player.tier_id.id)
            if tl:
                already_sold = env['auction.auction.player'].sudo().search_count([
                    ('auction_id', '=', auction.id),
                    ('player_id.tier_id', '=', player.tier_id.id),
                ])
                if already_sold >= tl[0].max_players:
                    return {
                        'success': False,
                        'error': 'Your team has reached the max players for the "%s" tier.' % player.tier_id.name,
                    }

        player.sudo().write({
            'current_bid': bid_amount,
            'current_bid_team_id': int(team_id),
        })
        player._restart_owner_bid_window()

        # Log this bid so it can be restored on a revoke
        env['auction.bid.log'].sudo().create({
            'player_id': player.id,
            'team_id': int(team_id),
            'bid_amount': bid_amount,
            'tournament_id': player.tournament_id.id,
        })

        return {
            'success': True,
            'current_bid': bid_amount,
            'team_name': auction.team_id.name,
        }

    # ── Revoke last bid ─────────────────────────────────────────────────────

    @http.route('/auction/owner/revoke-bid', type='json', auth='user', website=False, csrf=False)
    def revoke_bid(self, player_id, team_id, **kw):
        """Undo the owner's most recent bid, restoring the previous bid state."""
        user = request.env.user.sudo()
        env = request.env

        if not user.auction_team_id or user.auction_team_id.id != int(team_id):
            return {'success': False, 'error': 'You can only revoke bids for your own team.'}

        player = env['auction.team.player'].sudo().browse(int(player_id))
        if not player.exists() or not player.is_on_stage:
            return {'success': False, 'error': 'Player is not currently on stage.'}
        if player.state != 'auction':
            return {'success': False, 'error': 'This player is closed. Waiting for the next player.'}

        if not player.current_bid_team_id or player.current_bid_team_id.id != int(team_id):
            return {'success': False, 'error': 'You are not the current leader — can only revoke your own bid.'}

        env.cr.execute(
            'SELECT id FROM auction_team_player WHERE id = %s FOR UPDATE',
            (player.id,),
        )
        player.invalidate_cache()
        if player._owner_bid_window_closed(player):
            return {
                'success': False,
                'error': 'Time is up. Only the last bid stays open.',
            }

        tournament = player.tournament_id
        if not tournament.revoke_enabled:
            return {'success': False, 'error': 'Bid revoke is not active for this tournament.'}
        if tournament.max_revokes == 0:
            return {'success': False, 'error': 'Bid revoke is disabled (max revokes set to 0).'}
        if tournament.revokes_used >= tournament.max_revokes:
            return {
                'success': False,
                'error': 'All %d revoke(s) for this tournament have been used.' % tournament.max_revokes,
            }

        # Find all bid log entries for this player, in placement order
        logs = env['auction.bid.log'].sudo().search(
            [('player_id', '=', player.id)], order='id asc'
        )
        if not logs:
            return {'success': False, 'error': 'No bid history found for this player.'}

        last_log = logs[-1]
        if last_log.team_id.id != int(team_id):
            return {'success': False, 'error': 'Your last bid is not the most recent — cannot revoke.'}

        prev_log = logs[-2] if len(logs) > 1 else None

        if prev_log:
            player.sudo().write({
                'current_bid': prev_log.bid_amount,
                'current_bid_team_id': prev_log.team_id.id,
            })
            player._restart_owner_bid_window()
            prev_team_name = prev_log.team_id.name
            new_bid = prev_log.bid_amount
        else:
            player.sudo().write({
                'current_bid': 0,
                'current_bid_team_id': False,
                'owner_bid_deadline': False,
            })
            prev_team_name = None
            new_bid = 0

        last_log.sudo().unlink()
        tournament.sudo().write({'revokes_used': tournament.revokes_used + 1})

        revokes_left = tournament.max_revokes - tournament.revokes_used
        return {
            'success': True,
            'new_bid': new_bid,
            'prev_team_name': prev_team_name,
            'revokes_remaining': revokes_left,
        }


# ── Excel export of owner logins (created through Tournament Settings) ─────────
class AuctionOwnerExport(http.Controller):

    @http.route('/auction_owner/owner_users/export/<int:tournament_id>',
                type='http', auth='user', methods=['GET'])
    def export_owner_users(self, tournament_id, **kw):
        import io
        import re
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
        from openpyxl.utils import get_column_letter

        tournament = request.env['auction.tournament'].browse(tournament_id).exists()
        if not tournament:
            return request.not_found()
        try:
            tournament._owner_users_check_access()
        except Exception:
            return request.make_response(
                'You are not allowed to export owner logins for this tournament.',
                [('Content-Type', 'text/plain; charset=utf-8')], status=403,
            )
        # Passwords are admin-only fields: read them with sudo after the access check.
        teams = tournament.sudo()._tournament_team_records().filtered(
            lambda t: t.owner_user_id and t.owner_wizard_created
        ).sorted(lambda t: (t.name or '').lower())
        if not teams:
            return request.make_response(
                'No owner users created from Tournament Settings were found.',
                [('Content-Type', 'text/plain; charset=utf-8')], status=404,
            )

        base_url = request.env['ir.config_parameter'].sudo().get_param('web.base.url') or ''
        login_url = (base_url.rstrip('/') + '/web/login') if base_url else ''

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Owner Logins'
        ws['A1'] = tournament.name or 'Tournament'
        ws['A1'].font = Font(size=14, bold=True)
        ws['A2'] = 'Owner login credentials — share each row only with its team owner.'
        ws['A2'].font = Font(italic=True, color='666666')
        headers = ['S.No', 'Team', 'Owner', 'Username', 'Password', 'Login URL']
        head_fill = PatternFill('solid', fgColor='0B1B3F')
        thin = Side(style='thin', color='C9CED8')
        for col, title in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col, value=title)
            cell.font = Font(bold=True, color='F2C94C')
            cell.fill = head_fill
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
        for idx, team in enumerate(teams, 1):
            row = [idx, team.name or '', team.manager or '',
                   team.owner_user_id.login or '', team.owner_password or '', login_url]
            for col, value in enumerate(row, 1):
                cell = ws.cell(row=4 + idx, column=col, value=value)
                cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
                cell.alignment = Alignment(vertical='center', horizontal='center' if col == 1 else 'left')
        widths = [7, 28, 26, 34, 18, 34]
        for col, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.freeze_panes = 'A5'

        buf = io.BytesIO()
        wb.save(buf)
        safe = re.sub(r'[^A-Za-z0-9_-]+', '_', tournament.name or 'tournament').strip('_') or 'tournament'
        filename = '%s_owner_logins.xlsx' % safe
        return request.make_response(
            buf.getvalue(),
            headers=[
                ('Content-Type', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),
                ('Content-Disposition', 'attachment; filename="%s"' % filename),
                ('Cache-Control', 'no-store'),
            ],
        )
